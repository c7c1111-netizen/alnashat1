"""تصدير قاعدتي بيانات المعلمين والبرامج (Excel / CSV).

ملف Excel المُصدَّر بنفس أعمدة الاستيراد، فيمكن تعديله ثم رفعه مرة أخرى.
"""
from __future__ import annotations

import csv
import io

from . import model
from .template_map import SLOT_LABELS

PROGRAM_COLS = [("name", "اسم البرنامج"), ("domain", "مجال النشاط"), ("goal", "الهدف"),
                ("description", "وصف مختصر"), ("target", "الفئة المستهدفة"), ("grades", "الصفوف المستهدفة"),
                ("teacher", "المعلم المسؤول"), ("sessions", "عدد الحصص"), ("start", "تاريخ البداية"),
                ("end", "تاريخ النهاية"), ("time", "وقت التنفيذ"), ("slot", "الحصة"), ("notes", "ملاحظات")]
TEACHER_COLS = [("name", "اسم المعلم"), ("subject", "مادة التدريس"), ("stage", "المرحلة"),
                ("grades", "الصفوف التي يدرسها"), ("subject_periods", "عدد حصص المادة"),
                ("rate", "نسبة 10%"), ("allowed", "الحصص الممكن الاستفادة منها"),
                ("programs", "البرامج المسندة"), ("program_sessions", "عدد حصص كل برنامج"),
                ("start", "تاريخ بداية التنفيذ"), ("remaining", "المتبقي بعد التنفيذ")]


def program_rows(plan):
    teachers = {t["id"]: t["name"] for t in plan["teachers"]}
    for p in plan["programs"]:
        r = dict(p)
        r["teacher"] = teachers.get(p.get("teacher_id"), "")
        r["slot"] = SLOT_LABELS.get(p.get("slot"), "")
        r["start"] = (p.get("start") or "") and p["start"] + "هـ"
        r["end"] = (p.get("end") or "") and p["end"] + "هـ"
        yield [r.get(k, "") for k, _ in PROGRAM_COLS]


def teacher_rows(plan):
    for ts in model.teacher_summary(plan):
        t = ts["teacher"]
        yield [t.get("name", ""), t.get("subject", ""),
               model.STAGES.get(t.get("stage"), {}).get("label", t.get("stage", "")),
               t.get("grades", ""), t.get("subject_periods", ""), "10%", ts["allowed"],
               "، ".join(r["program"] for r in ts["rows"]),
               "، ".join(str(r["sessions"]) for r in ts["rows"]),
               ts["rows"][0]["start"] if ts["rows"] else "", ts["remaining"]]


def to_xlsx(plan) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    sheets = [("البرامج", PROGRAM_COLS, list(program_rows(plan))),
              ("المعلمون", TEACHER_COLS, list(teacher_rows(plan)))]
    for i, (title, cols, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.sheet_view.rightToLeft = True
        ws.append([c[1] for c in cols])
        for r in rows:
            ws.append(r)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F4C5C")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def programs_csv(plan) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c[1] for c in PROGRAM_COLS])
    for r in program_rows(plan):
        w.writerow(r)
    return buf.getvalue().encode("utf-8-sig")


def teachers_csv(plan) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c[1] for c in TEACHER_COLS])
    for r in teacher_rows(plan):
        w.writerow(r)
    return buf.getvalue().encode("utf-8-sig")
