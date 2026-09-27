"""استيراد البرامج من ملف خارجي: Excel / CSV / Word / PDF.

قواعد صارمة:
- الملف المرفوع هو المرجع الأساسي للبرامج والتواريخ.
- لا تُخترع برامج، ولا تُعدل أسماء البرامج (تُنقل حرفيًا)، ولا تُغير التواريخ.
- أي تعارض أو نقص يُعرض للمستخدم قبل التطبيق، ولا يُطبق شيء إلا بعد موافقته.
"""
from __future__ import annotations

import csv
import io
import os
import re
import unicodedata

from . import hijri
from .model import ALL_DOMAINS, new_id
from .template_map import SLOT_KEYS, SLOT_LABELS

# مرادفات عناوين الأعمدة
HEADERS = {
    "name": ["اسم البرنامج", "البرنامج", "البرامج", "اسم النشاط", "النشاط", "program", "name"],
    "domain": ["المجال", "مجال النشاط", "المجالات", "domain"],
    "goal": ["الهدف", "الأهداف", "goal"],
    "description": ["الوصف", "وصف مختصر", "وصف البرنامج", "description"],
    "target": ["الفئة المستهدفة", "المستهدفون", "المستهدف", "target"],
    "grades": ["الصفوف المستهدفة", "الصفوف", "الصف", "grades", "grade"],
    "teacher": ["المعلم المسؤول", "اسم المعلم", "المعلم", "المنفذ", "teacher"],
    "sessions": ["عدد الحصص", "الحصص", "عدد حصص البرنامج", "sessions"],
    "date": ["تاريخ التنفيذ", "التاريخ", "date"],
    "start": ["تاريخ البداية", "تاريخ بداية التنفيذ", "من", "start"],
    "end": ["تاريخ النهاية", "تاريخ نهاية التنفيذ", "إلى", "الى", "end"],
    "time": ["وقت التنفيذ", "الوقت", "time"],
    "slot": ["الحصة", "الفترة", "slot"],
    "notes": ["ملاحظات", "الملاحظات", "notes"],
}


def _n(s) -> str:
    s = unicodedata.normalize("NFKC", str(s or ""))
    s = s.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ة", "ه").replace("ى", "ي")
    return re.sub(r"[\s‏‎:_\-]+", "", s).lower()


_HEADER_INDEX = {}
for key, names in HEADERS.items():
    for nm in names:
        _HEADER_INDEX.setdefault(_n(nm), key)


def _match_header(cell) -> str | None:
    t = _n(cell)
    if not t:
        return None
    if t in _HEADER_INDEX:
        return _HEADER_INDEX[t]
    # تطابق جزئي للعناوين الطويلة
    best = None
    for nm, key in _HEADER_INDEX.items():
        if key == "name" and not t.startswith(_n("اسم البرنامج")):
            continue  # عمود اسم البرنامج يجب أن يطابق بدقة
        if len(nm) >= 4 and nm in t and (best is None or len(nm) > len(best[0])):
            best = (nm, key)
    return best[1] if best else None


def _find_header(rows) -> tuple[int, dict] | None:
    for i, row in enumerate(rows[:15]):
        cols = {}
        for j, c in enumerate(row):
            k = _match_header(c)
            if k and k not in cols:
                cols[k] = j
        if "name" in cols and len(cols) >= 2:
            return i, cols
    return None


def _fix_pdf_text(s: str) -> str:
    """توحيد نص خلية مستخرجة من PDF (ترتيب منطقي) وضم أسطرها."""
    return repair_pdf_arabic(" ".join(unicodedata.normalize("NFKC", ln).strip() for ln in (s or "").split("\n") if ln.strip()))


_LIGATURE_FIXES = [  # تسلسلات لا توجد في العربية الصحيحة، سببها ترميز «لا» المركبة في PDF
    (re.compile(r"ا([أإآ])ل"), r"ال\1"),
    (re.compile(r"(^|\s)([أإآ])ال"), r"\1ال\2"),
    (re.compile(r"(^|\s)اال"), r"\1الا"),
]


def repair_pdf_arabic(s: str) -> str:
    for rx, rep in _LIGATURE_FIXES:
        s = rx.sub(rep, s)
    return s


def _visual_to_logical(s: str) -> str:
    """نصوص PDF العربية كثيرًا ما تُستخرج بالترتيب البصري (معكوسة):
    نعكس كل سطر مع إبقاء الأرقام والنصوص اللاتينية بترتيبها، ثم نوحد الحروف
    (بعد العكس حتى تبقى «لا» المركبة صحيحة)."""
    out = []
    for ln in (s or "").split("\n"):
        r = ln[::-1]
        r = re.sub(r"[0-9A-Za-z/.\-:]+", lambda m: m.group(0)[::-1], r)
        out.append(unicodedata.normalize("NFKC", r).strip())
    return repair_pdf_arabic(" ".join(x for x in out if x))


# ----------------------------------------------------------------------------
# قراءة الجداول من كل نوع ملف
# ----------------------------------------------------------------------------
def _rows_csv(data: bytes):
    for enc in ("utf-8-sig", "cp1256", "utf-16"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("تعذر قراءة ترميز ملف CSV")
    dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t") if text.strip() else csv.excel
    return [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]


def _rows_xlsx(data: bytes):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    tables = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else (v if hasattr(v, "year") else str(v).strip()) for v in r])
        tables.append(rows)
    return tables


def _rows_docx(data: bytes):
    import zipfile

    from lxml import etree

    from .ooxml import NS, q

    z = zipfile.ZipFile(io.BytesIO(data))
    root = etree.fromstring(z.read("word/document.xml"))
    tables = []
    for tbl in root.iter(q("w:tbl")):
        rows = []
        for tr in tbl.findall("w:tr", NS):
            row = []
            for tc in tr.findall("w:tc", NS):
                txt = " ".join("".join(t.text or "" for t in p.iter(q("w:t"))) for p in tc.findall("w:p", NS)).strip()
                span = tc.find("w:tcPr/w:gridSpan", NS)
                row.append(txt)
                for _ in range(int(span.get(q("w:val"))) - 1 if span is not None else 0):
                    row.append(txt)
            rows.append(row)
        tables.append(rows)
    return tables


def _rows_pdf(data: bytes):
    import pymupdf

    tables = []
    doc = pymupdf.open(stream=data, filetype="pdf")
    for page in doc:
        for t in page.find_tables().tables:
            raw = [[c or "" for c in row] for row in t.extract()]
            rows = [[_fix_pdf_text(c) for c in row] for row in raw]
            if _find_header(rows) is None:
                vis = [[_visual_to_logical(c) for c in row] for row in raw]
                if _find_header(vis) is not None:
                    rows = vis
            tables.append(rows)
    return tables


def read_tables(filename: str, data: bytes):
    ext = os.path.splitext(filename.lower())[1]
    if ext in (".csv", ".txt", ".tsv"):
        return [_rows_csv(data)]
    if ext in (".xlsx", ".xlsm"):
        return _rows_xlsx(data)
    if ext == ".docx":
        return _rows_docx(data)
    if ext == ".pdf":
        return _rows_pdf(data)
    raise ValueError("نوع الملف غير مدعوم. الأنواع المدعومة: Excel (xlsx) و CSV و Word (docx) و PDF")


# ----------------------------------------------------------------------------
# التحليل وكشف التعارضات
# ----------------------------------------------------------------------------
def _date(v):
    if v in (None, ""):
        return "", None
    if hasattr(v, "year"):
        h = hijri.from_gregorian(v)
        return hijri.fmt(h, suffix=False), None
    try:
        h = hijri.parse(v)
        return hijri.fmt(h, suffix=False), None
    except hijri.DateError as e:
        return str(v), str(e)


def _slot(v):
    t = _n(v)
    if not t:
        return ""
    for key in SLOT_KEYS:
        if _n(SLOT_LABELS[key]) in t or t in _n(SLOT_LABELS[key]):
            return key
    m = re.search(r"\d", str(v))
    if m and 1 <= int(m.group()) <= 7:
        return f"p{m.group()}"
    return ""


def parse_programs(filename: str, data: bytes, plan: dict, structure: dict) -> dict:
    """يعيد {programs, issues, source}. لا يعدل الخطة."""
    tables = read_tables(filename, data)
    issues, programs = [], []
    teachers = {_n(t["name"]): t for t in plan.get("teachers", []) if t.get("name")}
    lo = hi = None
    try:
        from .model import semester_range
        lo, hi = semester_range(plan, structure)
    except Exception:  # noqa: BLE001
        pass
    found_table = False
    for ti, rows in enumerate(tables):
        hdr = _find_header(rows)
        if not hdr:
            continue
        found_table = True
        hi_row, cols = hdr
        for ri, row in enumerate(rows[hi_row + 1:], start=hi_row + 2):
            def get(k):
                j = cols.get(k)
                return row[j] if j is not None and j < len(row) else ""
            name = str(get("name") or "").strip()
            if not name and not any(str(c).strip() for c in row):
                continue
            where = f"جدول {ti + 1} - سطر {ri}"
            p = {"id": new_id("prg"), "name": name, "domain": str(get("domain") or "").strip(),
                 "goal": str(get("goal") or "").strip(), "description": str(get("description") or "").strip(),
                 "target": str(get("target") or "").strip(), "grades": str(get("grades") or "").strip(),
                 "time": str(get("time") or "").strip(), "notes": str(get("notes") or "").strip(),
                 "slot": _slot(get("slot")), "teacher_id": "", "mode": "sessions",
                 "weekdays": [0, 1, 2, 3, 4], "date_required": True, "source": f"{filename} ({where})",
                 "raw": {k: str(get(k)) for k in cols}}
            if not name:
                issues.append({"level": "error", "where": where, "msg": "سطر بدون اسم برنامج."})
            sess = str(get("sessions") or "").strip()
            if sess:
                m = re.search(r"\d+", sess.translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")))
                if m:
                    p["sessions"] = int(m.group())
                else:
                    p["sessions"] = ""
                    issues.append({"level": "error", "where": where, "msg": f"«{name}»: عدد الحصص «{sess}» غير مفهوم."})
            else:
                p["sessions"] = ""
            start_raw = get("start") or get("date")
            p["start"], err = _date(start_raw)
            if err:
                issues.append({"level": "error", "where": where, "msg": f"«{name}»: {err}"})
            p["end"], err = _date(get("end"))
            if err:
                issues.append({"level": "error", "where": where, "msg": f"«{name}»: {err}"})
            if not p["start"]:
                issues.append({"level": "warning", "where": where, "msg": f"«{name}»: لا يوجد تاريخ تنفيذ في الملف."})
            elif lo and hijri.try_parse(p["start"]):
                o = hijri.ordinal(hijri.parse(p["start"]))
                if not (lo <= o <= hi):
                    issues.append({"level": "warning", "where": where,
                                   "msg": f"«{name}»: التاريخ {p['start']}هـ خارج فترة الخطة الحالية."})
            if hasattr(start_raw, "year") or (str(start_raw) and re.search(r"(19|20)\d\d", str(start_raw))):
                issues.append({"level": "info", "where": where,
                               "msg": f"«{name}»: التاريخ في الملف ميلادي ({start_raw}) وحُوّل داخليًا إلى {p['start']}هـ."})
            if p["domain"] and p["domain"] not in ALL_DOMAINS:
                match = [d for d in ALL_DOMAINS if _n(d) == _n(p["domain"])]
                if match:
                    p["domain"] = match[0]
                else:
                    issues.append({"level": "warning", "where": where,
                                   "msg": f"«{name}»: المجال «{p['domain']}» غير موجود في القالب (لن يظهر في جدول المجالات)."})
            tname = str(get("teacher") or "").strip()
            if tname:
                t = teachers.get(_n(tname))
                if t:
                    p["teacher_id"] = t["id"]
                else:
                    p["teacher_name_new"] = tname
                    issues.append({"level": "warning", "where": where,
                                   "msg": f"«{name}»: المعلم «{tname}» غير موجود في قائمة المعلمين وسيُضاف عند التطبيق."})
            if p["slot"] in ("morning", "prayer", "routine"):
                p["mode"] = "daily"
            programs.append(p)
    if filename.lower().endswith(".pdf") and programs:
        issues.append({"level": "warning", "where": filename,
                       "msg": "الأسماء مستخرجة من PDF: راجع كتابتها قبل التطبيق، فبعض ملفات PDF تُخزّن الحروف العربية "
                              "بترتيب بصري أو بحروف مركبة. لأدق نتيجة استخدم ملف Excel أو Word."})
    if not found_table:
        issues.append({"level": "error", "where": filename,
                       "msg": "لم يُعثر على جدول يحتوي عمود «اسم البرنامج». تأكد من وجود صف عناوين في الملف."})

    # تعارضات داخل الملف نفسه
    by_name = {}
    for p in programs:
        if p["name"]:
            by_name.setdefault(_n(p["name"]), []).append(p)
    for group in by_name.values():
        if len(group) > 1:
            variants = {(g["start"], g["end"], str(g["sessions"]), g["domain"]) for g in group}
            if len(variants) > 1:
                issues.append({"level": "conflict", "where": "الملف",
                               "msg": f"البرنامج «{group[0]['name']}» مكرر ببيانات مختلفة: " +
                               " | ".join(f"التاريخ {v[0] or '-'} الحصص {v[2] or '-'}" for v in variants)})
            else:
                issues.append({"level": "info", "where": "الملف", "msg": f"البرنامج «{group[0]['name']}» مكرر بنفس البيانات."})
    # تعارضات مع البرامج الموجودة في الخطة
    existing = {_n(p["name"]): p for p in plan.get("programs", []) if p.get("name")}
    for key, group in by_name.items():
        ex = existing.get(key)
        for g in (group if ex else []):
            diffs = []
            if g["start"] and ex.get("start") and g["start"] != ex.get("start"):
                diffs.append(f"التاريخ في الخطة {ex['start']} وفي الملف {g['start']}")
            if str(g["sessions"]) and str(ex.get("sessions", "")) and str(g["sessions"]) != str(ex.get("sessions")):
                diffs.append(f"الحصص في الخطة {ex['sessions']} وفي الملف {g['sessions']}")
            if diffs:
                issues.append({"level": "conflict", "where": "الخطة الحالية",
                               "msg": f"«{g['name']}» موجود في الخطة ببيانات مختلفة: " + "، ".join(diffs) +
                               ". عند التطبيق تُعتمد بيانات الملف."})
    return {"programs": programs, "issues": issues, "source": filename}


def apply_import(plan: dict, programs: list, mode: str = "merge") -> dict:
    """تطبيق نتيجة الاستيراد بعد موافقة المستخدم.

    mode = "replace": البرامج المستوردة تحل محل كل برامج الخطة.
    mode = "merge": تُضاف البرامج، والموجود بنفس الاسم يُحدَّث ببيانات الملف.
    """
    teachers = {_n(t["name"]): t for t in plan["teachers"]}
    for p in programs:
        nm = p.pop("teacher_name_new", None)
        p.pop("raw", None)
        if nm:
            t = teachers.get(_n(nm))
            if not t:
                t = {"id": new_id("tch"), "name": nm, "subject": "", "stage": plan["school"].get("stage", ""),
                     "grades": "", "subject_periods": "", "notes": "أضيف من ملف الاستيراد"}
                plan["teachers"].append(t)
                teachers[_n(nm)] = t
            p["teacher_id"] = t["id"]
    if mode == "replace":
        kept_ids = set()
        plan["programs"] = programs
    else:
        idx = {_n(p["name"]): i for i, p in enumerate(plan["programs"]) if p.get("name")}
        for p in programs:
            i = idx.get(_n(p["name"]))
            if i is not None:
                p["id"] = plan["programs"][i]["id"]
                plan["programs"][i] = {**plan["programs"][i], **{k: v for k, v in p.items() if v not in ("", None)}}
            else:
                plan["programs"].append(p)
        kept_ids = {p["id"] for p in plan["programs"]}
    ids = {p["id"] for p in plan["programs"]}
    plan["schedule"] = [e for e in plan["schedule"] if e.get("program_id") in ids and (mode != "replace" or e.get("program_id") in kept_ids)]
    return plan
